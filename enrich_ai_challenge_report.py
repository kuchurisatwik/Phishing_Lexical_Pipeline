from __future__ import annotations

import asyncio
import ipaddress
from pathlib import Path

import aiohttp
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from extract.evidence_generator import (
    _detect_sector,
    _infer_cse_from_url,
    _load_cse_helpers,
    _lookup_cse_by_domain,
    take_screenshots,
)
from extract.feature_extractor import lookup_dns, normalize_hostname
from extract.rdap_whois import RDAPClient


INPUT = Path("ai_challenge.xlsx")
OUTPUT = Path("ai_challenge_final_format_enriched.xlsx")
EVIDENCE_DIR = Path("output") / "Evidence"
DNS_CONCURRENCY = 80
RDAP_CONCURRENCY = 80
IP_RDAP_CONCURRENCY = 60
CAPTURE_EVIDENCE = False


def na(value) -> str:
    if value is None:
        return "NA"
    text = str(value).strip()
    if text.lower() in {"", "nan", "nat", "none", "null", "-1"}:
        return "NA"
    return text


def clean_url(value) -> str:
    text = na(value)
    if text == "NA":
        return ""
    return text.replace("http:s//", "https://").replace("https:s//", "https://")


def phishing_label(value) -> str:
    text = str(value).strip().lower()
    if text in {"phishing", "yes", "confirmed_phishing"}:
        return "Yes"
    if text in {"suspected", "suspicious_needs_review"}:
        return "Suspected"
    return "No"


def domain_for_lookup(row: pd.Series) -> str:
    domain = normalize_hostname(row.get("domain"))
    if domain:
        return domain
    return normalize_hostname(clean_url(row.get("url")))


def report_cse(row: pd.Series, helpers: dict) -> str:
    for value in (row.get("domain"), clean_url(row.get("url"))):
        cse_name = _lookup_cse_by_domain(value, helpers)
        if cse_name != "NA":
            return cse_name
    return _infer_cse_from_url(clean_url(row.get("url")), helpers)


def extract_ip_org_country(data: dict) -> tuple[str, str]:
    org = data.get("name") or data.get("handle") or ""
    country = data.get("country") or ""

    for entity in data.get("entities", []) or []:
        roles = {str(role).lower() for role in entity.get("roles", [])}
        if roles and not ({"registrant", "administrative", "technical", "abuse"} & roles):
            continue

        if not country:
            country = entity.get("country") or ""

        vcard = entity.get("vcardArray", [])
        if len(vcard) >= 2:
            for entry in vcard[1]:
                if len(entry) >= 4 and entry[0] in {"org", "fn"} and entry[3] and not org:
                    org = str(entry[3])
                if len(entry) >= 4 and entry[0] == "adr" and entry[3] and not country:
                    address = entry[3]
                    if isinstance(address, list) and address:
                        country = str(address[-1] or "")

        if org and country:
            break

    return na(org), na(country)


async def load_ip_bootstrap(session: aiohttp.ClientSession, version: int) -> list[tuple[ipaddress._BaseNetwork, str]]:
    mapping: list[tuple[ipaddress._BaseNetwork, str]] = []
    url = f"https://data.iana.org/rdap/ipv{version}.json"
    try:
        async with session.get(url, timeout=aiohttp.ClientTimeout(total=15)) as resp:
            if resp.status != 200:
                return mapping
            data = await resp.json(content_type=None)
    except Exception:
        return mapping

    for networks, urls in data.get("services", []):
        if not urls:
            continue
        base = urls[0].rstrip("/")
        for network in networks:
            try:
                mapping.append((ipaddress.ip_network(network, strict=False), base))
            except Exception:
                pass
    return mapping


def find_ip_base(ip: str, v4map: list, v6map: list) -> str:
    try:
        address = ipaddress.ip_address(ip)
    except Exception:
        return ""
    mapping = v4map if address.version == 4 else v6map
    for network, base in mapping:
        if address in network:
            return base
    return ""


async def lookup_ip_rdap(
    session: aiohttp.ClientSession,
    ip: str,
    v4map: list,
    v6map: list,
    sem: asyncio.Semaphore,
) -> tuple[str, str]:
    if not ip:
        return "NA", "NA"

    base = find_ip_base(ip, v4map, v6map) or "https://rdap.org"
    url = f"{base}/ip/{ip}"
    try:
        async with sem:
            async with session.get(
                url,
                timeout=aiohttp.ClientTimeout(total=10),
                headers={"Accept": "application/rdap+json"},
            ) as resp:
                if resp.status != 200:
                    return "NA", "NA"
                data = await resp.json(content_type=None)
    except Exception:
        return "NA", "NA"

    return extract_ip_org_country(data)


def style_workbook(path: Path) -> None:
    wb = load_workbook(path)
    ws = wb.active
    ws.title = "final_format_enriched"

    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = {
        "A": 42,
        "B": 44,
        "C": 24,
        "D": 36,
        "E": 18,
        "F": 28,
        "G": 20,
        "H": 38,
        "I": 28,
        "J": 22,
        "K": 70,
        "L": 16,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    wb.save(path)


async def main() -> None:
    df = pd.read_excel(INPUT, sheet_name="model _predictions")
    helpers = _load_cse_helpers()
    df["__lookup_domain"] = df.apply(domain_for_lookup, axis=1)
    domains = sorted({domain for domain in df["__lookup_domain"] if domain})

    dns_sem = asyncio.Semaphore(DNS_CONCURRENCY)
    dns_done = 0

    async def dns_lookup(domain: str):
        nonlocal dns_done
        async with dns_sem:
            result = await asyncio.to_thread(lookup_dns, domain)
            dns_done += 1
            if dns_done % 100 == 0 or dns_done == len(domains):
                print(f"DNS {dns_done}/{len(domains)}", flush=True)
            return domain, result

    dns_pairs = await asyncio.gather(*(dns_lookup(domain) for domain in domains))
    dns_cache = dict(dns_pairs)

    async with aiohttp.ClientSession() as session:
        rdap = RDAPClient(
            session,
            rdap_concurrency=RDAP_CONCURRENCY,
            whois_concurrency=1,
            rdap_timeout=8,
            whois_timeout=1,
            whois_delay=0.2,
            rdap_retries=1,
        )
        await rdap.init()
        reg_sem = asyncio.Semaphore(RDAP_CONCURRENCY)
        rdap_done = 0

        async def reg_lookup(domain: str):
            nonlocal rdap_done
            async with reg_sem:
                # Use RDAP only here. The normal helper falls back to WHOIS, which
                # is intentionally slow and can make challenge-scale reports hang.
                result = await rdap._rdap_lookup(domain)
                rdap_done += 1
                if rdap_done % 100 == 0 or rdap_done == len(domains):
                    print(f"RDAP {rdap_done}/{len(domains)}", flush=True)
                return result

        reg_results = await asyncio.gather(*(reg_lookup(domain) for domain in domains))
        reg_cache = dict(zip(domains, reg_results))
        rdap.shutdown()

        ips = sorted(
            {
                ip.strip()
                for result in dns_cache.values()
                for ip in (result.resolved_ips or "").split(";")
                if ip.strip()
            }
        )
        v4map, v6map = await asyncio.gather(
            load_ip_bootstrap(session, 4),
            load_ip_bootstrap(session, 6),
        )
        ip_sem = asyncio.Semaphore(IP_RDAP_CONCURRENCY)
        ip_results = await asyncio.gather(
            *(lookup_ip_rdap(session, ip, v4map, v6map, ip_sem) for ip in ips)
        )
        ip_cache = dict(zip(ips, ip_results))

    phishing_urls = []
    for _, row in df.iterrows():
        if phishing_label(row.get("Phishng(yes)")) == "Yes":
            phishing_urls.append((clean_url(row.get("url")) or str(row.get("url")), report_cse(row, helpers)))

    screenshot_map = {}
    if phishing_urls and CAPTURE_EVIDENCE:
        try:
            screenshot_map = await asyncio.to_thread(take_screenshots, phishing_urls, str(EVIDENCE_DIR))
        except Exception as exc:
            print(f"Evidence capture skipped: {type(exc).__name__}: {exc}", flush=True)

    rows = []
    for _, row in df.iterrows():
        domain = row["__lookup_domain"]
        dns = dns_cache.get(domain)
        info = reg_cache.get(domain)
        ips = dns.resolved_ips if dns else ""
        first_ip = next((ip.strip() for ip in ips.split(";") if ip.strip()), "")
        hosting_isp, hosting_country = ip_cache.get(first_ip, ("NA", "NA"))
        cse_name = report_cse(row, helpers)
        prediction = na(row.get("Phishng(yes)"))
        mapped_sector = na(row.get("mapped_cse_sector"))
        cleaned_url = clean_url(row.get("url"))

        rows.append(
            {
                "Identified Domain Name": str(row.get("url", "")).strip(),
                "Corresponding CSE Name": cse_name,
                "IP Address": na(ips),
                "Hosting ISP": hosting_isp,
                "Hosting Country": hosting_country,
                "Registrant Name": na(getattr(info, "registrant_name", "")),
                "Registrant Country": "NA",
                "Name Servers": na(";".join(getattr(info, "nameservers", ()) or ())),
                "Evidence File Path": screenshot_map.get(cleaned_url, "NA"),
                "Source of Detection": _detect_sector(cse_name),
                "Remarks": f"AI challenge prediction: {prediction}; mapped sector: {mapped_sector}",
                "Phishing (Yes)": phishing_label(prediction),
            }
        )

    report_df = pd.DataFrame(rows)
    report_df.to_excel(OUTPUT, index=False, engine="openpyxl")
    style_workbook(OUTPUT)

    print(OUTPUT.resolve())
    print("rows", len(report_df))
    for column in [
        "IP Address",
        "Hosting ISP",
        "Hosting Country",
        "Registrant Name",
        "Name Servers",
        "Evidence File Path",
    ]:
        print(column, int((report_df[column].astype(str) != "NA").sum()))
    print(report_df["Phishing (Yes)"].value_counts().to_string())


if __name__ == "__main__":
    asyncio.run(main())
